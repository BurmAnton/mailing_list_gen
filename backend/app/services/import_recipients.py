import csv
import io
import json
from typing import Any

from email_validator import EmailNotValidError, validate_email
from openpyxl import load_workbook

KNOWN_FIELDS = {"name", "industry", "description", "website_url"}
COMPANY_ALIASES = {"company_name", "company", "компания", "companyname", "org", "organization"}
CELL_MAX = 500
ROW_MAX = 3000


def _normalize_header(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "_")


def _cell_str(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _truncate(text: str, limit: int) -> str:
    if len(text) <= limit:
        return text
    return text[: limit - 1] + "…"


def _build_row_from_mapped(mapped: dict[str, Any], original_headers: list[str] | None = None) -> dict[str, Any]:
    """Build recipient row: email required; all other columns -> extra_fields."""
    email_raw = _cell_str(mapped.get("email"))
    if not email_raw:
        return {"__invalid__": True}

    try:
        validated = validate_email(email_raw, check_deliverability=False)
        email = validated.normalized
    except EmailNotValidError:
        return {"__invalid__": True, "email": email_raw}

    extra: dict[str, str] = {}
    known: dict[str, str | None] = {k: None for k in KNOWN_FIELDS}
    company_name: str | None = None

    # Preserve column order from original headers when available
    keys = original_headers or list(mapped.keys())
    for key in keys:
        norm = _normalize_header(key)
        if not norm or norm == "email":
            continue
        val = _cell_str(mapped.get(norm) if norm in mapped else mapped.get(key))
        if val is None:
            continue
        val = _truncate(val, CELL_MAX)

        if norm in COMPANY_ALIASES and not company_name:
            company_name = val
        if norm in KNOWN_FIELDS:
            known[norm] = val
        # Always keep in extra_fields (except pure email)
        display_key = str(key).strip() if not original_headers else str(key).strip()
        # Use normalized key for stability
        extra[norm] = val

    if not company_name:
        for key, val in extra.items():
            if key not in KNOWN_FIELDS and val:
                company_name = val
                break
    if not company_name:
        company_name = "—"

    # Cap total extra payload size
    serialized = json.dumps(extra, ensure_ascii=False)
    if len(serialized) > ROW_MAX:
        # Drop trailing keys until under limit
        trimmed: dict[str, str] = {}
        for k, v in extra.items():
            candidate = {**trimmed, k: v}
            if len(json.dumps(candidate, ensure_ascii=False)) > ROW_MAX:
                break
            trimmed[k] = v
        extra = trimmed

    return {
        "email": email,
        "company_name": company_name,
        "name": known.get("name"),
        "industry": known.get("industry"),
        "description": known.get("description"),
        "website_url": known.get("website_url"),
        "extra_fields": extra,
        "source_row": json.dumps(
            {k: (_cell_str(v) if not isinstance(v, str) else v) for k, v in mapped.items()},
            ensure_ascii=False,
        ),
    }


def parse_csv_bytes(content: bytes) -> list[dict[str, Any]]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise ValueError("CSV пустой или без заголовков")
    headers = [_normalize_header(h) for h in reader.fieldnames]
    if "email" not in headers:
        raise ValueError("В файле нет обязательной колонки: email")

    rows: list[dict[str, Any]] = []
    for raw in reader:
        mapped = {_normalize_header(k): v for k, v in raw.items()}
        rows.append(_build_row_from_mapped(mapped, headers))
    return rows


def parse_xlsx_bytes(content: bytes) -> list[dict[str, Any]]:
    wb = load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise ValueError("XLSX пустой") from exc

    headers = [_normalize_header(h) for h in header_row]
    if "email" not in headers:
        raise ValueError("В файле нет обязательной колонки: email")

    rows: list[dict[str, Any]] = []
    for values in rows_iter:
        mapped = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
        if all(v is None or str(v).strip() == "" for v in mapped.values()):
            continue
        rows.append(_build_row_from_mapped(mapped, headers))
    return rows


def parse_recipients_file(filename: str, content: bytes) -> tuple[list[dict[str, Any]], int]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        parsed = parse_csv_bytes(content)
    elif lower.endswith(".xlsx") or lower.endswith(".xls"):
        parsed = parse_xlsx_bytes(content)
    else:
        raise ValueError("Поддерживаются только CSV и XLSX")

    valid_rows: list[dict[str, Any]] = []
    invalid = 0
    for row in parsed:
        if row.get("__invalid__"):
            invalid += 1
            continue
        valid_rows.append(row)
    return valid_rows, invalid
