import hashlib
import logging
import re
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.campaign import Campaign
from app.models.rag_file import CampaignRagFile
from app.services.llm_client import LLMClient

logger = logging.getLogger(__name__)

RAG_ROOT = Path(__file__).resolve().parents[2] / "data" / "rag"
DIGEST_MAX_CHARS = 6000
RAW_COMBINED_LIMIT = 40000
HEURISTIC_DIGEST_THRESHOLD = 6000

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv", ".xlsx", ".xls", ".pptx"}


def _extract_xlsx_text(content: bytes) -> str:
    import io
    from openpyxl import load_workbook

    wb = load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _extract_pptx_text(content: bytes) -> str:
    import io
    from pptx import Presentation

    prs = Presentation(io.BytesIO(content))
    lines: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text and shape.text.strip():
                lines.append(shape.text.strip())
    return "\n".join(lines)


def extract_text_from_bytes(filename: str, content: bytes) -> str:
    lower = filename.lower()
    suffix = Path(lower).suffix
    if suffix == ".doc":
        raise ValueError(
            "Старый формат Word (.doc) не поддерживается. Сохраните файл как .docx и загрузите снова."
        )
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError(
            f"Неподдерживаемый тип файла: {suffix or '(без расширения)'}. "
            f"Разрешены: {', '.join(sorted(ALLOWED_EXTENSIONS))} (не .doc)."
        )

    if not content:
        raise ValueError("Файл пустой")

    if suffix in {".txt", ".md", ".csv"}:
        for encoding in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="ignore")

    if suffix == ".pdf":
        from pypdf import PdfReader
        import io

        try:
            reader = PdfReader(io.BytesIO(content))
            parts = []
            for page in reader.pages:
                parts.append(page.extract_text() or "")
            text = "\n".join(parts)
            if not text.strip():
                raise ValueError(
                    "PDF без извлекаемого текста (возможно скан). "
                    "Загрузите DOCX/TXT или добавьте сайт через URL."
                )
            return text
        except ValueError as exc:
            msg = str(exc)
            if "извлекаемого текста" in msg or "DOCX/TXT" in msg:
                raise
            raise ValueError(
                f"Не удалось прочитать PDF ({exc}). "
                "Проверьте, что файл не повреждён и содержит текст."
            ) from exc
        except Exception as exc:
            raise ValueError(
                f"Не удалось прочитать PDF ({exc}). "
                "Проверьте, что файл не повреждён и содержит текст."
            ) from exc

    if suffix == ".pptx":
        try:
            text = _extract_pptx_text(content)
            if not text.strip():
                raise ValueError("PPTX не содержит текста на слайдах")
            return text
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"Не удалось прочитать PPTX: {exc}") from exc

    if suffix in {".xlsx", ".xls"}:
        try:
            if suffix == ".xls":
                import xlrd

                book = xlrd.open_workbook(file_contents=content)
                lines: list[str] = []
                for sheet in book.sheets():
                    for rx in range(sheet.nrows):
                        cells = [
                            str(sheet.cell_value(rx, cx)).strip()
                            for cx in range(sheet.ncols)
                            if str(sheet.cell_value(rx, cx)).strip()
                        ]
                        if cells:
                            lines.append(" | ".join(cells))
                text = "\n".join(lines)
            else:
                text = _extract_xlsx_text(content)
            if not text.strip():
                raise ValueError("Excel не содержит текста в ячейках")
            return text
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"Не удалось прочитать Excel ({suffix}): {exc}") from exc

    if suffix == ".docx":
        import io
        from docx import Document

        doc = Document(io.BytesIO(content))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    raise ValueError(f"Не удалось извлечь текст из {filename}")


def ensure_rag_dir(campaign_id: int) -> Path:
    path = RAG_ROOT / str(campaign_id)
    path.mkdir(parents=True, exist_ok=True)
    return path


def _clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def heuristic_digest(combined: str, max_chars: int = DIGEST_MAX_CHARS) -> str:
    if len(combined) <= max_chars:
        return combined
    # Keep head + tail for coverage
    head = max_chars * 2 // 3
    tail = max_chars - head - 20
    return combined[:head] + "\n\n[...]\n\n" + combined[-tail:]


def compute_files_hash(files: list[CampaignRagFile]) -> str:
    h = hashlib.sha256()
    for f in sorted(files, key=lambda x: x.id or 0):
        h.update(str(f.id).encode())
        h.update((f.filename or "").encode())
        h.update((f.extracted_text or "").encode())
    return h.hexdigest()


async def save_rag_file(
    session: AsyncSession,
    *,
    campaign_id: int,
    filename: str,
    content: bytes,
    content_type: str | None,
) -> CampaignRagFile:
    text = _clean_text(extract_text_from_bytes(filename, content))
    folder = ensure_rag_dir(campaign_id)
    safe_name = re.sub(r"[^\w.\-]+", "_", filename)[:180]
    stored = folder / f"{safe_name}"
    # Avoid overwrite collisions
    if stored.exists():
        stem = stored.stem
        suffix = stored.suffix
        i = 1
        while stored.exists():
            stored = folder / f"{stem}_{i}{suffix}"
            i += 1
    stored.write_bytes(content)

    rag = CampaignRagFile(
        campaign_id=campaign_id,
        filename=filename,
        stored_path=str(stored),
        content_type=content_type,
        source_type="file",
        source_url=None,
        extracted_text=text,
        char_count=len(text),
    )
    session.add(rag)
    await _invalidate_digest(session, campaign_id)
    await session.flush()
    return rag


async def save_rag_website(
    session: AsyncSession,
    *,
    campaign_id: int,
    url: str,
) -> CampaignRagFile:
    from urllib.parse import urlparse

    from app.services.enrichment import fetch_website_text

    parsed = urlparse(url.strip())
    if parsed.scheme not in ("http", "https") or not parsed.netloc:
        raise ValueError("Укажите корректный URL (http/https)")

    normalized = url.strip()
    text = _clean_text(await fetch_website_text(normalized))
    if not text:
        raise ValueError("Не удалось извлечь текст с сайта")

    # Cap website RAG text a bit higher than per-company scrape default
    text = text[:12000]
    host = parsed.netloc
    path = (parsed.path or "/").rstrip("/") or "/"
    filename = f"site:{host}{path}"[:500]

    rag = CampaignRagFile(
        campaign_id=campaign_id,
        filename=filename,
        stored_path="",
        content_type="text/html",
        source_type="website",
        source_url=normalized,
        extracted_text=text,
        char_count=len(text),
    )
    session.add(rag)
    await _invalidate_digest(session, campaign_id)
    await session.flush()
    return rag


async def _invalidate_digest(session: AsyncSession, campaign_id: int) -> None:
    campaign = await session.get(Campaign, campaign_id)
    if campaign:
        campaign.rag_digest_text = None
        campaign.rag_digest_hash = None


async def delete_rag_file(session: AsyncSession, rag: CampaignRagFile) -> None:
    if rag.stored_path:
        path = Path(rag.stored_path)
        if path.exists():
            try:
                path.unlink()
            except OSError as exc:
                logger.warning("Failed to delete rag file %s: %s", path, exc)
    await _invalidate_digest(session, rag.campaign_id)
    await session.delete(rag)


async def build_or_get_rag_digest(
    session: AsyncSession,
    campaign: Campaign,
    client: LLMClient | None = None,
) -> str:
    result = await session.execute(
        select(CampaignRagFile).where(CampaignRagFile.campaign_id == campaign.id).order_by(CampaignRagFile.id)
    )
    files = list(result.scalars().all())
    if not files:
        campaign.rag_digest_text = ""
        campaign.rag_digest_hash = ""
        await session.flush()
        return ""

    file_hash = compute_files_hash(files)
    if campaign.rag_digest_hash == file_hash and campaign.rag_digest_text is not None:
        return campaign.rag_digest_text

    parts: list[str] = []
    total = 0
    for f in files:
        chunk = (f.extracted_text or "").strip()
        if not chunk:
            continue
        header = f"### {f.filename}\n"
        if getattr(f, "source_type", "file") == "website" and f.source_url:
            header = f"### website: {f.source_url}\n"
        remaining = RAW_COMBINED_LIMIT - total
        if remaining <= 0:
            break
        body = chunk[:remaining]
        parts.append(header + body)
        total += len(header) + len(body)

    combined = "\n\n".join(parts)
    if not combined.strip():
        digest = ""
    elif len(combined) <= HEURISTIC_DIGEST_THRESHOLD:
        digest = heuristic_digest(combined)
    else:
        # Compress with LLM once for the campaign
        try:
            llm = client or LLMClient(model=campaign.llm_model, provider=campaign.llm_provider)
            messages = [
                {
                    "role": "system",
                    "content": (
                        "Сжми материалы для RAG-контекста персонализированных B2B-писем. "
                        "Сохрани факты, оффер, УТП, ограничения, CTA, тон. "
                        f"Ответ обычным текстом, не более {DIGEST_MAX_CHARS} символов."
                    ),
                },
                {"role": "user", "content": combined[:RAW_COMBINED_LIMIT]},
            ]
            digest = await llm.chat_completion(
                messages,
                temperature=0.2,
                enable_web_search=False,
            )
            digest = heuristic_digest(_clean_text(digest), DIGEST_MAX_CHARS)
        except Exception as exc:
            logger.warning("RAG digest LLM failed, using heuristic: %s", exc)
            digest = heuristic_digest(combined)

    campaign.rag_digest_text = digest
    campaign.rag_digest_hash = file_hash
    await session.flush()
    return digest
